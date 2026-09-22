from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName


OUTPUT_DIR = Path(__file__).parent
YELLOW = PatternFill(fill_type="solid", fgColor="FFFF00")


def add_data_sheet(workbook, candidate):
    sheet = workbook.create_sheet("Data")
    sheet.append(
        [
            "Measure_Cd",
            "Unit",
            "Measure_Value",
            "Comment",
            "Organisation_Cd",
            "Business_Unit_Cd",
            "Observation_Period_Cd",
            "Status",
        ]
    )
    rows = [
        ["APR3F_06", "Number", 100, "Good", "AFW", "AddCtrl", "1980", "Open"],
        ["APR3F_06", "Number", 125, "Conflicting duplicate", "AFW", "AddCtrl", "1980", "Open"],
        ["APR3G_11", "Text", 125 if candidate else "confirmed", "Value ¬¬ confirmed", "AFW", "Bio", "1980 APR", "Unknown"],
        ["B0407HP_NRP", "Number", "TBC", "FORBIDDEN", "AFW", "Bio", "1980", "Closed"],
        ["B0407HP_RP", "Number", None, "", "AFW", "Bio", "1980", "Open"],
    ]
    for row in rows:
        sheet.append(row)

    # Cells in the yellow input area exercise required-filled and
    # unexpected-formula checks.
    sheet["Z11"].fill = YELLOW
    sheet["Z11"] = None if candidate else "organisation-code"
    sheet["Z12"].fill = YELLOW
    sheet["Z12"] = "=1+1" if candidate else "user input"

    sheet["A7"] = "=1+3" if candidate else "=2+2"

    return sheet


def add_codes_sheet(workbook, candidate):
    sheet = workbook.create_sheet("Codes")
    sheet.append(["Code"])
    for code in [
        "APR3F_06",
        "APR3G_11",
        "B0407HP_NRP",
        "B0407HP_RP",
    ]:
        sheet.append([code])
    return sheet


def add_formula_sheet(workbook, candidate):
    sheet = workbook.create_sheet("Formula_Checks")
    sheet["A1"] = "Formula"
    sheet["A2"] = "=1/0" if candidate else "=1+1"
    if candidate:
        sheet["A3"] = "#VALUE!"
        sheet["A3"].data_type = "e"
    else:
        sheet["A3"] = "=2+2"
    sheet["A4"] = "=1+1" if candidate else "=1+1"
    return sheet


def add_presentation_sheet(workbook):
    sheet = workbook.create_sheet("Presentation")
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Merged heading"
    sheet.append(["Chart", 1])
    chart = BarChart()
    chart.title = "Example chart"
    chart.add_data(Reference(sheet, min_col=2, min_row=2, max_row=2))
    sheet.add_chart(chart, "E2")
    return sheet


def create_workbook(path, candidate):
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_data_sheet(workbook, candidate)
    add_codes_sheet(workbook, candidate)
    add_formula_sheet(workbook, candidate)
    add_presentation_sheet(workbook)
    workbook.create_sheet("Guidance")
    if not candidate:
        workbook.create_sheet("ReferenceOnly")
    workbook.defined_names.add(
        DefinedName("InputCell", attr_text="'Data'!$Z$11")
    )
    workbook.save(path)


OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
create_workbook(OUTPUT_DIR / "candidate.xlsx", candidate=True)
create_workbook(OUTPUT_DIR / "reference.xlsx", candidate=False)
